from openpyxl import Workbook
from openpyxl import load_workbook

EXCEL_FILE_NAME = 'Database.xlsx'
db = load_workbook(filename=EXCEL_FILE_NAME)

seminar_db  = db['Seminar']
lecture_db  = db['Lecture']
user_db     = db['User_data']
answer_db   = db['Answer']

def get_user_row(user_key):
    for idx, row in enumerate(user_db.rows):
        if idx != 0 and row[0].value == user_key:
            user = row
            return user

        if idx == user_db.max_row - 1:
            user_db[user_db.max_row+1][0].value = user_key
            user_db[user_db.max_row][6].value = 0
            db.save(EXCEL_FILE_NAME)
            user = user_db[user_db.max_row]
            return user

def get_response(content):
    data = []
    for row in answer_db:
        if row[1].value == content:
            for i in range(2, 7):
               data.append(row[i].value)

    response = {
        "message" : {
            data[0] : data[1]
        },
        "keyboard" : {
            "type" : data[2],
            data[2] : data[3].split(",")
        }
    }
    return response
